"""
document_parser.py
-------------------
Extracts plain text from résumé files (PDF, DOCX, TXT).

Design notes:
- PDF extraction uses `pdfplumber` (pure-Python, no system dependencies like
  poppler required), which handles most text-based PDFs well.
- DOCX extraction uses `python-docx`, reading paragraphs and table cells
  (résumés frequently use tables for layout).
- Both functions accept either a file path (str/Path) or a file-like object
  (e.g. Streamlit's UploadedFile, or an in-memory BytesIO), so the same code
  works whether called from the CLI, tests, or the Streamlit app.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Union

import docx
import pdfplumber

FileLike = Union[str, Path, io.BytesIO, "io.IOBase"]


class DocumentParsingError(Exception):
    """Raised when a document cannot be parsed or contains no extractable text."""


def extract_text(file: FileLike, filename: str | None = None) -> str:
    """
    Dispatch to the right extractor based on file extension.

    Args:
        file: path to the file, or a file-like/bytes object (e.g. Streamlit upload).
        filename: original filename, required when `file` is not a path
                  (used to determine the extension).

    Returns:
        Cleaned plain text extracted from the document.

    Raises:
        DocumentParsingError: unsupported format, corrupt file, or no text found.
    """
    name = filename or (str(file) if isinstance(file, (str, Path)) else "")
    ext = Path(name).suffix.lower()

    if ext == ".pdf":
        raw = _extract_from_pdf(file)
    elif ext == ".docx":
        raw = _extract_from_docx(file)
    elif ext == ".doc":
        raise DocumentParsingError(
            "Le format .doc (Word 97-2003) n'est pas pris en charge. "
            "Merci de convertir le fichier en .docx ou .pdf."
        )
    elif ext == ".txt":
        raw = _extract_from_txt(file)
    else:
        raise DocumentParsingError(
            f"Format de fichier non pris en charge : '{ext or 'inconnu'}'. "
            "Formats acceptés : PDF, DOCX, TXT."
        )

    return _clean_text(raw)


def _extract_from_pdf(file: FileLike) -> str:
    try:
        with pdfplumber.open(file) as pdf:
            pages_text = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                pages_text.append(text)
            return "\n".join(pages_text)
    except Exception as exc:  # noqa: BLE001 - surface as a domain error
        raise DocumentParsingError(f"Impossible de lire le PDF : {exc}") from exc


def _extract_from_docx(file: FileLike) -> str:
    try:
        document = docx.Document(file)
        parts = [p.text for p in document.paragraphs]

        # Résumés often use tables for layout (skills grids, date/role columns)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)

        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        raise DocumentParsingError(f"Impossible de lire le DOCX : {exc}") from exc


def _extract_from_txt(file: FileLike) -> str:
    try:
        if isinstance(file, (str, Path)):
            return Path(file).read_text(encoding="utf-8", errors="ignore")
        data = file.read()
        if isinstance(data, bytes):
            return data.decode("utf-8", errors="ignore")
        return data
    except Exception as exc:  # noqa: BLE001
        raise DocumentParsingError(f"Impossible de lire le fichier texte : {exc}") from exc


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse excessive blank lines
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    text = text.strip()

    if not text:
        raise DocumentParsingError(
            "Aucun texte n'a pu être extrait de ce document (il s'agit peut-être "
            "d'un scan image sans couche de texte, ou d'un fichier vide)."
        )
    return text
